"""Document text extraction for PDF, DOCX and TXT inputs.

The public entry point is :func:`extract_text`, which accepts a path, raw bytes
or a file-like object (such as a Streamlit ``UploadedFile``) and returns clean
plain text.

Design rule: this module never raises a bare exception at the caller. Anything
that goes wrong is converted into :class:`DocumentParsingError` with a message
that is safe and useful to show to an end user.
"""

from __future__ import annotations

import io
import logging
import os
from pathlib import Path
from typing import Any

from .config import (
    MAX_FILE_SIZE_MB,
    MAX_TEXT_CHARS,
    MIN_USABLE_TEXT_CHARS,
    SUPPORTED_EXTENSIONS,
)

LOGGER = logging.getLogger(__name__)


class DocumentParsingError(Exception):
    """Raised when a document cannot be turned into usable text."""


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #


def _read_bytes(source: Any) -> tuple[bytes, str]:
    """Return ``(payload, filename)`` for a path / bytes / file-like object."""
    if isinstance(source, (str, os.PathLike)):
        path = Path(source)
        if not path.exists():
            raise DocumentParsingError(f"File not found: {path.name}")
        if not path.is_file():
            raise DocumentParsingError(f"Not a file: {path.name}")
        return path.read_bytes(), path.name

    if isinstance(source, (bytes, bytearray)):
        return bytes(source), ""

    # File-like (Streamlit UploadedFile, open() handle, BytesIO...)
    name = getattr(source, "name", "") or ""
    try:
        if hasattr(source, "seek"):
            source.seek(0)
        payload = source.read()
    except Exception as exc:  # pragma: no cover - defensive
        raise DocumentParsingError("The uploaded file could not be read.") from exc

    if isinstance(payload, str):
        payload = payload.encode("utf-8", errors="replace")
    if not isinstance(payload, (bytes, bytearray)):
        raise DocumentParsingError("Unsupported input type for text extraction.")
    return bytes(payload), os.path.basename(str(name))


def _check_size(payload: bytes, filename: str) -> None:
    size_mb = len(payload) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise DocumentParsingError(
            f"{filename or 'This file'} is {size_mb:.1f} MB. "
            f"The limit is {MAX_FILE_SIZE_MB:.0f} MB - please upload a smaller file."
        )


def detect_file_type(filename: str, payload: bytes | None = None) -> str:
    """Detect the document type from the extension, falling back to magic bytes.

    Returns one of ``"pdf"``, ``"docx"``, ``"txt"``.
    Raises :class:`DocumentParsingError` for anything else.
    """
    ext = Path(filename or "").suffix.lower()

    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in (".txt", ".md"):
        return "txt"

    if payload:
        if payload[:5] == b"%PDF-":
            return "pdf"
        if payload[:2] == b"PK":  # docx is a zip container
            return "docx"
        try:
            payload[:2048].decode("utf-8")
            return "txt"
        except UnicodeDecodeError:
            pass

    if ext == ".doc":
        raise DocumentParsingError(
            "Legacy .doc files are not supported. "
            "Please save the document as .docx or .pdf and upload again."
        )
    supported = ", ".join(SUPPORTED_EXTENSIONS)
    raise DocumentParsingError(
        f"Unsupported file type '{ext or 'unknown'}'. Supported formats: {supported}."
    )


def clean_extracted_text(text: str) -> str:
    """Normalise whitespace and strip artefacts left behind by PDF extraction."""
    if not text:
        return ""

    text = text.replace("\x00", " ").replace("\ufeff", "")
    # Ligatures and typographic characters that break skill matching.
    replacements = {
        "\ufb00": "ff",
        "\ufb01": "fi",
        "\ufb02": "fl",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u00a0": " ",
        "\u2022": "\n- ",
        "\u25cf": "\n- ",
        "\u25aa": "\n- ",
        "\u00ad": "",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)

    lines = [line.rstrip() for line in text.splitlines()]
    out: list[str] = []
    blank_run = 0
    for line in lines:
        stripped = " ".join(line.split())
        if not stripped:
            blank_run += 1
            if blank_run <= 1:
                out.append("")
            continue
        blank_run = 0
        out.append(stripped)

    cleaned = "\n".join(out).strip()
    if len(cleaned) > MAX_TEXT_CHARS:
        LOGGER.warning("Text truncated from %d to %d chars", len(cleaned), MAX_TEXT_CHARS)
        cleaned = cleaned[:MAX_TEXT_CHARS]
    return cleaned


# --------------------------------------------------------------------------- #
# Format-specific extractors
# --------------------------------------------------------------------------- #


def extract_text_from_pdf(source: Any) -> str:
    """Extract text from a PDF using PyMuPDF."""
    payload, filename = _read_bytes(source)
    _check_size(payload, filename)

    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24
    except ImportError:
        try:
            import fitz  # older PyMuPDF releases
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise DocumentParsingError(
                "PDF support requires PyMuPDF. Install it with: pip install pymupdf"
            ) from exc

    pages: list[str] = []
    try:
        with fitz.open(stream=payload, filetype="pdf") as doc:
            if doc.needs_pass:
                raise DocumentParsingError(
                    "This PDF is password protected. "
                    "Please remove the password or upload a TXT/DOCX version."
                )
            for page in doc:
                pages.append(page.get_text("text"))
    except DocumentParsingError:
        raise
    except Exception as exc:
        LOGGER.exception("PDF extraction failed")
        raise DocumentParsingError(
            "This PDF could not be opened. It may be corrupted - "
            "try re-saving or exporting it again."
        ) from exc

    text = clean_extracted_text("\n".join(pages))
    if len(text) < MIN_USABLE_TEXT_CHARS:
        raise DocumentParsingError(
            "Almost no text could be read from this PDF. It is probably a scanned "
            "image. Please upload a text-based PDF, a DOCX, or paste the text directly."
        )
    return text


def extract_text_from_docx(source: Any) -> str:
    """Extract text from a DOCX file, including tables."""
    payload, filename = _read_bytes(source)
    _check_size(payload, filename)

    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentParsingError(
            "DOCX support requires python-docx. Install it with: pip install python-docx"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(payload))
    except Exception as exc:
        LOGGER.exception("DOCX extraction failed")
        raise DocumentParsingError(
            "This DOCX file could not be opened. It may be corrupted or may "
            "actually be an older .doc file saved with a .docx extension."
        ) from exc

    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))

    text = clean_extracted_text("\n".join(parts))
    if len(text) < MIN_USABLE_TEXT_CHARS:
        raise DocumentParsingError(
            "This DOCX file appears to be empty. Please check the document and try again."
        )
    return text


def extract_text_from_txt(source: Any) -> str:
    """Extract text from a plain-text file, trying a few common encodings."""
    payload, filename = _read_bytes(source)
    _check_size(payload, filename)

    text = ""
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = payload.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:  # pragma: no cover - latin-1 practically never fails
        text = payload.decode("utf-8", errors="replace")

    text = clean_extracted_text(text)
    if len(text) < MIN_USABLE_TEXT_CHARS:
        raise DocumentParsingError(
            "This file appears to be empty. Please upload a file that contains text."
        )
    return text


# --------------------------------------------------------------------------- #
# Public entry point
# --------------------------------------------------------------------------- #

_EXTRACTORS = {
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "txt": extract_text_from_txt,
}


def extract_text(source: Any, filename: str | None = None) -> str:
    """Extract plain text from a resume or job description document.

    Args:
        source: A filesystem path, raw ``bytes``, or a file-like object.
        filename: Optional name used for type detection when ``source`` is
            bytes or an unnamed stream.

    Returns:
        Cleaned plain text.

    Raises:
        DocumentParsingError: With a user-facing message for every failure mode
            (unsupported type, corrupted file, empty file, scanned PDF, ...).
    """
    payload, detected_name = _read_bytes(source)
    name = filename or detected_name

    if not payload:
        raise DocumentParsingError("The file is empty (0 bytes). Please upload a valid document.")

    _check_size(payload, name)
    file_type = detect_file_type(name, payload)
    return _EXTRACTORS[file_type](payload)


def extract_text_safe(source: Any, filename: str | None = None) -> tuple[str, str | None]:
    """Non-raising variant of :func:`extract_text`.

    Returns:
        ``(text, error_message)``. Exactly one of the two is meaningful:
        on success ``error_message`` is ``None``; on failure ``text`` is ``""``.
    """
    try:
        return extract_text(source, filename), None
    except DocumentParsingError as exc:
        return "", str(exc)
    except Exception as exc:  # pragma: no cover - last-resort guard
        LOGGER.exception("Unexpected extraction failure")
        return "", f"Unexpected problem while reading the file: {type(exc).__name__}"


def read_text_file(path: str | os.PathLike[str]) -> str:
    """Convenience helper used by tests and the evaluation script."""
    return extract_text(Path(path))


__all__ = [
    "DocumentParsingError",
    "clean_extracted_text",
    "detect_file_type",
    "extract_text",
    "extract_text_from_docx",
    "extract_text_from_pdf",
    "extract_text_from_txt",
    "extract_text_safe",
    "read_text_file",
]
